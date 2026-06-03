from datetime import datetime
from io import BytesIO
from typing import Any, Dict, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas


def _safe_get_sentiment_count(batch: Dict[str, Any], key: str) -> int:
    sentiment_counts = batch.get("sentiment_counts", {})
    return int(sentiment_counts.get(key, 0))


def generate_excel_report(batch: Dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Batch Summary"

    summary_sheet["A1"] = "Sentrix Analysis Report"
    summary_sheet["A1"].font = Font(size=16, bold=True)
    summary_sheet.merge_cells("A1:B1")

    summary_sheet["A3"] = "Batch ID"
    summary_sheet["B3"] = str(batch.get("batch_id", ""))
    summary_sheet["A4"] = "Generated At"
    created_at = batch.get("created_at")
    if isinstance(created_at, datetime):
        summary_sheet["B4"] = created_at.isoformat()
    else:
        summary_sheet["B4"] = str(created_at or "")

    summary_sheet["A5"] = "Total Comments"
    summary_sheet["B5"] = int(batch.get("total_comments", 0))
    summary_sheet["A6"] = "Positive"
    summary_sheet["B6"] = _safe_get_sentiment_count(batch, "positive")
    summary_sheet["A7"] = "Negative"
    summary_sheet["B7"] = _safe_get_sentiment_count(batch, "negative")
    summary_sheet["A8"] = "Neutral"
    summary_sheet["B8"] = _safe_get_sentiment_count(batch, "neutral")

    details_sheet = workbook.create_sheet(title="Comment Analysis")
    headers = [
        "Comment ID",
        "Stakeholder Type",
        "Provision No",
        "Date",
        "Original Comment",
        "Cleaned Comment",
        "Language",
        "Sentiment",
        "Positive Score",
        "Neutral Score",
        "Negative Score",
        "Summary",
    ]
    details_sheet.append(headers)

    for cell in details_sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    results: Iterable[Dict[str, Any]] = batch.get("results", [])
    for result in results:
        confidence = result.get("sentiment_confidence", {})
        details_sheet.append(
            [
                result.get("comment_id", ""),
                result.get("stakeholder_type", ""),
                result.get("provision_no", ""),
                result.get("date", ""),
                result.get("original_comment", ""),
                result.get("cleaned_comment", ""),
                result.get("language", ""),
                result.get("sentiment", ""),
                float(confidence.get("positive", 0.0)),
                float(confidence.get("neutral", 0.0)),
                float(confidence.get("negative", 0.0)),
                result.get("summary", ""),
            ]
        )

    for column in ("A", "B", "C", "D", "G", "H"):
        details_sheet.column_dimensions[column].width = 18
    details_sheet.column_dimensions["E"].width = 60
    details_sheet.column_dimensions["F"].width = 60
    details_sheet.column_dimensions["L"].width = 70

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_pdf_report(batch: Dict[str, Any]) -> bytes:
    page_width, page_height = A4
    margin = 40
    line_height = 14
    max_width = page_width - (margin * 2)
    batch_id = str(batch.get("batch_id", ""))

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)

    def start_page() -> float:
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(margin, page_height - margin, "Sentrix Analysis Report")
        pdf.setFont("Helvetica", 9)
        pdf.drawRightString(
            page_width - margin, page_height - margin, f"Batch ID: {batch_id}"
        )
        pdf.line(
            margin,
            page_height - margin - 6,
            page_width - margin,
            page_height - margin - 6,
        )
        return page_height - margin - 24

    def draw_wrapped_block(
        y_pos: float, label: str, value: Any, value_font: str = "Helvetica"
    ) -> float:
        value_text = str(value or "")
        label_text = f"{label}: "

        pdf.setFont("Helvetica-Bold", 10)
        label_width = pdf.stringWidth(label_text, "Helvetica-Bold", 10)
        value_width = max_width - label_width
        wrapped_lines = simpleSplit(value_text, value_font, 10, value_width)
        if not wrapped_lines:
            wrapped_lines = [""]

        if y_pos < margin + 30:
            pdf.showPage()
            y_pos = start_page()

        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(margin, y_pos, label_text)
        pdf.setFont(value_font, 10)
        pdf.drawString(margin + label_width, y_pos, wrapped_lines[0])
        y_pos -= line_height

        for line in wrapped_lines[1:]:
            if y_pos < margin + 20:
                pdf.showPage()
                y_pos = start_page()
                pdf.setFont(value_font, 10)
            pdf.drawString(margin, y_pos, line)
            y_pos -= line_height

        return y_pos

    y = start_page()
    created_at = batch.get("created_at")
    created_at_text = (
        created_at.isoformat() if isinstance(created_at, datetime) else str(created_at or "")
    )

    y = draw_wrapped_block(y, "Generated At", created_at_text)
    y = draw_wrapped_block(y, "Total Comments", batch.get("total_comments", 0))
    y = draw_wrapped_block(y, "Positive", _safe_get_sentiment_count(batch, "positive"))
    y = draw_wrapped_block(y, "Negative", _safe_get_sentiment_count(batch, "negative"))
    y = draw_wrapped_block(y, "Neutral", _safe_get_sentiment_count(batch, "neutral"))

    y -= 8
    if y < margin + 30:
        pdf.showPage()
        y = start_page()
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(margin, y, "Comment-wise Analysis")
    y -= line_height

    results: Iterable[Dict[str, Any]] = batch.get("results", [])
    for index, result in enumerate(results, start=1):
        if y < margin + 80:
            pdf.showPage()
            y = start_page()

        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(margin, y, f"{index}. Comment ID: {result.get('comment_id', '')}")
        y -= line_height

        y = draw_wrapped_block(y, "Stakeholder", result.get("stakeholder_type", ""))
        y = draw_wrapped_block(y, "Provision", result.get("provision_no", ""))
        y = draw_wrapped_block(y, "Date", result.get("date", ""))
        y = draw_wrapped_block(y, "Sentiment", result.get("sentiment", ""))
        y = draw_wrapped_block(y, "Summary", result.get("summary", ""))
        y = draw_wrapped_block(y, "Original Comment", result.get("original_comment", ""))

        if y < margin + 20:
            pdf.showPage()
            y = start_page()
        else:
            pdf.line(margin, y, page_width - margin, y)
            y -= 10

    pdf.save()
    return buffer.getvalue()
